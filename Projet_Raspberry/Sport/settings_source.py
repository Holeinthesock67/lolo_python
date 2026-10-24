"""Lieu de la météo, lu dans l'onglet « Réglages » de la feuille « Coach – Carnet ».

Chaque matin, le Pi télécharge la feuille en .xlsx (rclone, accès Google en lecture seule limité au dossier
« Coach - Garmin export »), lit la ligne « Lieu météo » et la convertit en coordonnées avec le géocodage
gratuit d'Open-Meteo. Valeurs acceptées : un nom de ville (« Sarre-Union »), une ville et un pays
(« Biarritz, France ») ou des coordonnées (« 43.48 ; -1.56 »).

Aucun lieu n'est écrit dans le code : si la lecture échoue, le dernier lieu connu (data/location.json) est
conservé ; à défaut, la section [location] de config.toml sert de secours.
"""
import json
import logging
import re
import subprocess
from datetime import datetime
from pathlib import Path

import requests

log = logging.getLogger("settings")
GEOCODE = "https://geocoding-api.open-meteo.com/v1/search"
DEFAULT_REMOTE = "coachro:Coach – Carnet.xlsx"
DEFAULT_KEY = "Lieu météo"
COORDS = re.compile(r"^\s*(-?\d{1,2}(?:[.,]\d+)?)\s*[;/]\s*(-?\d{1,3}(?:[.,]\d+)?)\s*$|"
                    r"^\s*(-?\d{1,2}(?:\.\d+)?)\s*,\s*(-?\d{1,3}(?:\.\d+)?)\s*$")


def fetch_sheet(remote, dest, timeout=180):
    """Télécharge la feuille (exportée en .xlsx par rclone). Renvoie le chemin, ou None en cas d'échec."""
    dest = Path(dest)
    tmp = dest.with_suffix(".new.xlsx")
    try:
        r = subprocess.run(["rclone", "copyto", "--contimeout", "30s", "--timeout", "60s", remote, str(tmp)],
                           capture_output=True, text=True, timeout=timeout)
    except (OSError, subprocess.TimeoutExpired) as e:
        log.warning("Réglages : feuille non téléchargée (%s)", e)
        return None
    if r.returncode != 0 or not tmp.exists():
        log.warning("Réglages : feuille non téléchargée (%s)", (r.stderr or "").strip()[-200:] or r.returncode)
        return None
    tmp.replace(dest)
    return dest


def read_setting(xlsx, key=DEFAULT_KEY, sheet="Réglages"):
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            return None
        for row in wb[sheet].iter_rows(values_only=True):
            if row and row[0] is not None and str(row[0]).strip().lower() == key.lower():
                v = row[1] if len(row) > 1 else None
                return str(v).strip() if v not in (None, "") and str(v).strip() else None
        return None
    finally:
        wb.close()


def geocode(query, get=requests.get):
    """« Ville » ou « Ville, pays » ou « lat ; lon » -> dict(name, latitude, longitude, timezone), ou None."""
    m = COORDS.match(query)
    if m:
        lat, lon = (float(x.replace(",", ".")) for x in (m.group(1, 2) if m.group(1) else m.group(3, 4)))
        if -90 <= lat <= 90 and -180 <= lon <= 180:
            return dict(name=query.strip(), latitude=lat, longitude=lon, timezone="auto")
        return None
    city, _, hint = (s.strip() for s in query.partition(","))
    r = get(GEOCODE, params=dict(name=city, count=10, language="fr", format="json"), timeout=30)
    r.raise_for_status()
    res = r.json().get("results") or []
    if hint:
        h = hint.lower()
        res = [x for x in res if h in " ".join(str(x.get(k, "")) for k in ("country", "country_code", "admin1",
                                                                          "admin2", "postcodes")).lower()] or []
    if not res:
        return None
    x = res[0]
    label = ", ".join(v for v in (x.get("name"), x.get("admin1"), x.get("country")) if v)
    return dict(name=label, latitude=x["latitude"], longitude=x["longitude"], timezone=x.get("timezone") or "auto")


def cached_place(data_dir):
    """Nom du dernier lieu utilisé (pour l'instantané quand la météo n'a pas été relancée)."""
    try:
        return json.loads((Path(data_dir) / "location.json").read_text(encoding="utf-8")).get("name")
    except (OSError, ValueError):
        return None


def resolve(cfg, data_dir, fetch=fetch_sheet, geocode_fn=geocode):
    """Lieu à utiliser pour la météo. Ne lève jamais d'exception : renvoie None si aucun lieu n'est connu."""
    s = cfg.get("settings", {})
    remote, key = s.get("remote", DEFAULT_REMOTE), s.get("key", DEFAULT_KEY)
    data_dir = Path(data_dir)
    data_dir.mkdir(parents=True, exist_ok=True)
    cache = data_dir / "location.json"
    try:
        cached = json.loads(cache.read_text(encoding="utf-8")) if cache.exists() else None
    except ValueError:
        cached = None

    query = None
    if remote:
        x = fetch(remote, data_dir / "carnet.xlsx")
        if x:
            try:
                query = read_setting(x, key)
            except Exception as e:   # fichier illisible, openpyxl absent…
                log.warning("Réglages illisibles : %s", e)
            if not query:
                log.warning("Réglages : ligne « %s » absente ou vide", key)
    if query:
        if cached and cached.get("query") == query:
            return cached
        try:
            loc = geocode_fn(query)
        except Exception as e:
            log.warning("Géocodage de « %s » impossible : %s", query, e)
            loc = None
        if loc:
            loc.update(query=query, updated=datetime.now().isoformat(timespec="seconds"))
            cache.write_text(json.dumps(loc, ensure_ascii=False, indent=1), encoding="utf-8")
            log.info("Lieu météo : %s (%.4f, %.4f)", loc["name"], loc["latitude"], loc["longitude"])
            return loc
        log.warning("Lieu « %s » introuvable : dernier lieu connu conservé", query)
    if cached:
        return cached
    loc = cfg.get("location")
    if loc and "latitude" in loc and "longitude" in loc:
        return dict(name=loc.get("name", "config.toml"), latitude=loc["latitude"], longitude=loc["longitude"],
                    timezone=loc.get("timezone", "auto"), query=None)
    return None
