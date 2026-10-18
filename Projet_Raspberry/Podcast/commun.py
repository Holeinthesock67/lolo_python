"""Outils partagés par veille.py et podcast.py.

- journal (log) et exécution de commandes (run) ;
- rclone avec délai maximal et nouvelles tentatives ;
- lecture de la configuration des émissions depuis le Google Sheet « Configuration podcasts »,
  téléchargé en .xlsx par rclone. Si le Drive est injoignable, la dernière copie est utilisée.
"""
import datetime as dt
import re
import subprocess
import sys
import time
import unicodedata
from pathlib import Path

HOME = Path.home()
DATA = HOME / "podcast-data"
REMOTE = "podcast:"                     # racine : dossier « Podcasts » sur le Drive
CONFIG_REMOTE = REMOTE + "Configuration podcasts.xlsx"
CONFIG_LOCAL = DATA / "configuration.xlsx"

JOURS = ["lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche"]
FORMATS = ["debat", "revue"]
DEFAUTS = {"age_max_jours": 3, "caracteres_par_contenu": 8000, "caracteres_au_total": 300000,
           "voix_a": "fr-FR-VivienneMultilingualNeural", "voix_b": "fr-FR-RemyMultilingualNeural",
           "prenom_a": "Camille", "prenom_b": "Thomas", "duree_min": 30}


# ---------- journal et commandes ----------

def log(msg):
    print(f"[{dt.datetime.now():%Y-%m-%d %H:%M}] {msg}", file=sys.stderr, flush=True)


def run(cmd, timeout=None):
    res = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
    if res.returncode != 0:
        raise RuntimeError(f"{cmd[0]} a échoué : {res.stderr.strip()[-300:]}")
    return res


RCLONE_TIMEOUT = 300   # durée maximale d'un transfert (secondes) avant abandon
RCLONE_TRIES = 3


def rclone(verb, src, dst):
    """rclone copy / copyto avec délai maximal et nouvelles tentatives.
    --contimeout / --timeout : rclone abandonne lui-même une connexion muette."""
    cmd = ["rclone", verb, "--contimeout", "30s", "--timeout", "60s", src, dst]
    for attempt in range(1, RCLONE_TRIES + 1):
        try:
            return run(cmd, timeout=RCLONE_TIMEOUT)
        except (subprocess.TimeoutExpired, RuntimeError) as exc:
            if "directory not found" in str(exc):   # dossier absent : réessayer ne sert à rien
                raise FileNotFoundError(f"{src} : dossier introuvable")
            reason = (f"bloqué plus de {RCLONE_TIMEOUT} s" if isinstance(exc, subprocess.TimeoutExpired)
                      else str(exc)[:200])
            if attempt == RCLONE_TRIES:
                raise RuntimeError(f"rclone {src} -> {dst} : échec après {RCLONE_TRIES} essais ({reason})")
            log(f"  rclone essai {attempt}/{RCLONE_TRIES} raté ({reason}), nouvel essai dans {30 * attempt} s")
            time.sleep(30 * attempt)


def rclone_copy(src, dst):
    return rclone("copy", src, dst)


def today_name():
    return JOURS[dt.date.today().weekday()]


# ---------- configuration (Google Sheet) ----------

def _key(header):
    """« Durée (min) » -> « duree_min », « Chaîne ou URL » -> « chaine_ou_url »."""
    s = unicodedata.normalize("NFKD", str(header or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")


def _text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _yes(v):
    return _text(v).lower() in ("oui", "o", "yes", "y", "true", "vrai", "1", "x")


def _int(v, default):
    try:
        return int(float(_text(v)))
    except ValueError:
        return default


def _rows(ws):
    """Lignes d'un onglet sous forme de dictionnaires {clé normalisée: valeur}."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    keys = [_key(h) for h in rows[0]]
    out = []
    for n, row in enumerate(rows[1:], start=2):
        d = {k: v for k, v in zip(keys, row) if k}
        if any(_text(v) for v in d.values()):
            d["_ligne"] = n
            out.append(d)
    return out


def _parse_source(row, tab, warnings):
    name = _text(row.get("nom"))
    where = _text(row.get("chaine_ou_url"))
    kind = _text(row.get("type")).lower()
    where_msg = f"{tab}, ligne {row['_ligne']} ({name or 'sans nom'})"
    if not name or not where:
        warnings.append(f"{where_msg} : nom ou chaîne/URL manquant, source ignorée")
        return None
    src = {"name": name, "type": kind, "lang": _text(row.get("langue")) or "fr",
           "camp": _text(row.get("camp")) or "?", "these": _text(row.get("these")),
           "filter": _yes(row.get("filtre_mots_cles")),
           "max_items": _int(row.get("contenus_examines"), 5),
           "enabled": _yes(row.get("active")) if _text(row.get("active")) else True}
    if kind == "youtube":
        p = re.search(r"[?&]list=(PL[\w-]+)", where) or re.fullmatch(r"(PL[\w-]{10,})", where)
        m = re.search(r"/channel/(UC[\w-]{22})", where) or re.fullmatch(r"(UC[\w-]{22})", where)
        if p:   # playlist, rangée des plus récentes aux plus anciennes
            src["playlist_id"] = p.group(1)
        elif m:
            src["channel_id"] = m.group(1)
        else:
            m = re.search(r"@([\w.-]+)", where)
            if not m:
                warnings.append(f"{where_msg} : chaîne YouTube illisible « {where} », source ignorée")
                return None
            src["handle"] = m.group(1)
    elif kind == "rss":
        if not where.startswith("http"):
            warnings.append(f"{where_msg} : l'URL du flux doit commencer par http, source ignorée")
            return None
        src["url"] = where
    else:
        warnings.append(f"{where_msg} : type « {kind} » inconnu (youtube ou rss), source ignorée")
        return None
    return src


def parse_config(path):
    """Lit le classeur et renvoie (émissions, avertissements)."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    tabs = {ws.title: ws for ws in wb.worksheets}
    warnings, emissions, ids = [], [], set()
    if "Émissions" not in tabs:
        raise RuntimeError("onglet « Émissions » introuvable dans le Sheet")
    for row in _rows(tabs["Émissions"]):
        where_msg = f"Émissions, ligne {row['_ligne']}"
        eid = _text(row.get("identifiant")).lower()
        if not re.fullmatch(r"[a-z0-9_-]+", eid or ""):
            warnings.append(f"{where_msg} : identifiant « {eid} » invalide (minuscules, chiffres, - ou _), ignorée")
            continue
        if eid in ids:
            warnings.append(f"{where_msg} : identifiant « {eid} » en double, ignorée")
            continue
        fmt = _text(row.get("format")).lower()
        jour = _text(row.get("jour")).lower()
        tab = _text(row.get("onglet_des_sources"))
        dossier = _text(row.get("dossier_drive"))
        problems = []
        if fmt not in FORMATS:
            problems.append(f"format « {fmt} » inconnu ({', '.join(FORMATS)})")
        if jour not in JOURS:
            problems.append(f"jour « {jour} » inconnu")
        if tab not in tabs:
            problems.append(f"onglet de sources « {tab} » introuvable")
        if not dossier:
            problems.append("dossier Drive manquant")
        if problems:
            warnings.append(f"{where_msg} ({eid}) : " + " ; ".join(problems) + ", ignorée")
            continue
        ids.add(eid)
        sources = [s for s in (_parse_source(r, tab, warnings) for r in _rows(tabs[tab])) if s]
        em = {"id": eid, "nom": _text(row.get("nom")) or eid, "actif": _yes(row.get("active")),
              "format": fmt, "jour": jour, "dossier": dossier, "onglet": tab,
              "consignes": _text(row.get("consignes")),
              "keywords": [k.strip() for k in _text(row.get("mots_cles")).split(",") if k.strip()],
              "sources": sources}
        for k in ("voix_a", "voix_b", "prenom_a", "prenom_b"):
            em[k] = _text(row.get(k)) or DEFAUTS[k]
        em["duree"] = _int(row.get("duree_min"), DEFAUTS["duree_min"])
        # clés attendues par veille.py
        em["max_age_days"] = _int(row.get("age_max_jours"), DEFAUTS["age_max_jours"])
        em["max_chars_per_item"] = _int(row.get("caracteres_par_contenu"), DEFAUTS["caracteres_par_contenu"])
        em["max_total_chars"] = _int(row.get("caracteres_au_total"), DEFAUTS["caracteres_au_total"])
        em["data_dir"] = str(DATA / eid)
        em["remote"] = f"{REMOTE}{dossier}"
        if any(s["filter"] for s in sources) and not em["keywords"]:
            warnings.append(f"{where_msg} ({eid}) : des sources filtrent par mots-clés mais aucun mot-clé n'est défini")
        emissions.append(em)
    return emissions, warnings


def load_emissions(refresh=True):
    """Télécharge le Sheet (sauf refresh=False) puis le lit. Garde la dernière copie en secours."""
    DATA.mkdir(parents=True, exist_ok=True)
    if refresh:
        tmp = CONFIG_LOCAL.with_suffix(".new.xlsx")
        try:
            rclone("copyto", CONFIG_REMOTE, str(tmp))
            parse_config(tmp)            # ne remplace la copie que si le fichier est lisible
            tmp.replace(CONFIG_LOCAL)
        except Exception as exc:
            tmp.unlink(missing_ok=True)
            if not CONFIG_LOCAL.exists():
                raise RuntimeError(f"configuration introuvable : {exc}")
            age = dt.datetime.now() - dt.datetime.fromtimestamp(CONFIG_LOCAL.stat().st_mtime)
            log(f"ATTENTION : Sheet injoignable ({str(exc)[:150]}). "
                f"Copie locale utilisée, vieille de {age.days} j {age.seconds // 3600} h.")
    return parse_config(CONFIG_LOCAL)


def select(emissions, wanted=None):
    """Émissions actives, ou celle demandée (même inactive)."""
    if wanted:
        found = [e for e in emissions if e["id"] == wanted]
        if not found:
            raise SystemExit(f"Émission « {wanted} » introuvable dans le Sheet.")
        return found
    return [e for e in emissions if e["actif"]]
